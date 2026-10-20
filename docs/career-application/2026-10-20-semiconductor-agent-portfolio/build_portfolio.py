"""Build portfolio.pdf, portfolio.docx and figures/roadmap.png from portfolio.md.

Usage: python build_portfolio.py --fonts DIR
DIR must contain NotoSansSC-Regular.ttf, NotoSansSC-Medium.ttf and NotoSansSC-Bold.ttf
(static instances of the OFL Noto Sans SC variable font at wght 400/500/700, each with its own
PostScript name; instances that keep the shared name are merged by PDF viewers and bold is lost).
Requires reportlab, python-docx and poppler's pdftoppm.
"""

import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Flowable, KeepTogether, ListFlowable, ListItem, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

HERE = Path(__file__).resolve().parent

# Palette taken from the author's resume: navy headings, steel-blue accents, gray meta text.
NAVY = "#1F2A44"
BLUE = "#2F5D8C"
TEXT = "#2B2F36"
GRAY = "#6B7280"
RULE = "#C9CED6"
TABLE_HEAD = "#EEF1F5"
TABLE_LINE = "#D5DAE1"
BOX_FILL = "#F5F7FA"

RUNNING_HEAD = "鄢澳勇 | 项目说明：半导体智能排程 Agent"


# ---------- markdown parsing ----------

def parse_blocks(md):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h", level, line[level:].strip()))
            i += 1
        elif line.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            blocks.append(("img", m.group(1), m.group(2)))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif re.match(r"(- |\d+\. )", line):
            ordered = bool(re.match(r"\d+\. ", line))
            items = []
            while i < len(lines) and re.match(r"(- |\d+\. )", lines[i]):
                items.append(re.sub(r"^(- |\d+\. )", "", lines[i]))
                i += 1
            blocks.append(("list", ordered, items))
        else:
            para = []
            while i < len(lines) and lines[i].strip() and not re.match(r"(#|\||!\[|- |\d+\. )", lines[i]):
                para.append(lines[i].strip())
                i += 1
            blocks.append(("p", " ".join(para)))
    return blocks


def inline_runs(text):
    """Split inline markdown into (text, kind, href) runs; kind in plain/bold/code/link."""
    runs = []
    for m in re.finditer(r"\*\*(.+?)\*\*|`(.+?)`|\[(.+?)\]\((.+?)\)|([^*`\[]+|[*`\[])", text):
        if m.group(1):
            runs.append((m.group(1), "bold", None))
        elif m.group(2):
            runs.append((m.group(2), "code", None))
        elif m.group(3):
            runs.append((m.group(3), "link", m.group(4)))
        else:
            runs.append((m.group(5), "plain", None))
    return runs


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_rl(text):
    out = []
    for s, kind, href in inline_runs(text):
        if kind == "bold":
            out.append(f'<font name="SC-Bold" color="{NAVY}">{esc(s)}</font>')
        elif kind == "code":
            out.append(f'<font color="{GRAY}">{esc(s)}</font>')
        elif kind == "link":
            out.append(f'<a href="{href}" color="{BLUE}"><u>{esc(s)}</u></a>')
        else:
            out.append(esc(s))
    return "".join(out)


# ---------- roadmap figure ----------

STAGES = [
    ("1  共同梳理", ["和工程师一起排产", "整理规则、特殊情况"], ["试点范围明确", "验收目标经工程师确认"]),
    ("2  离线回放", ["打通 Harness 最小路径", "用案例集回放对比"], ["不违反硬约束", "差异可解释、中断可恢复"]),
    ("3  只读影子运行", ["和现场并行只出建议", "开展培训，记录纠正"], ["工程师能核对结果", "异常时能正确接管"]),
    ("4  限定场景使用", ["人工审核后执行", "按新问题更新规则"], ["业务价值与成本达标", "维护责任清楚"]),
]


def roadmap_drawing(width=170 * mm):
    h = 62 * mm
    d = Drawing(width, h)
    d.add(String(0, h - 5 * mm, "四个阶段逐级过门，达不到条件就停在当前阶段",
                 fontName="SC-Bold", fontSize=9.5, fillColor=colors.HexColor(NAVY)))
    gap = 6 * mm
    bw = (width - 3 * gap) / 4
    bh = 20 * mm
    top = h - 9 * mm
    for k, (title, body, gate) in enumerate(STAGES):
        x = k * (bw + gap)
        first = k == 0
        d.add(Rect(x, top - bh, bw, bh, rx=2.5, ry=2.5,
                   fillColor=colors.HexColor("#E6EDF5" if first else BOX_FILL),
                   strokeColor=colors.HexColor(BLUE if first else RULE), strokeWidth=1.1 if first else 0.8))
        d.add(String(x + bw / 2, top - 6 * mm, title, fontName="SC-Bold", fontSize=9,
                     fillColor=colors.HexColor(NAVY), textAnchor="middle"))
        for j, t in enumerate(body):
            d.add(String(x + bw / 2, top - 11.5 * mm - j * 4.2 * mm, t, fontName="SC", fontSize=7.4,
                         fillColor=colors.HexColor(TEXT), textAnchor="middle"))
        if k < 3:
            ax, ay = x + bw + 1 * mm, top - bh / 2
            d.add(Line(ax, ay, ax + gap - 2.6 * mm, ay, strokeColor=colors.HexColor(GRAY), strokeWidth=0.8))
            d.add(Polygon([ax + gap - 2 * mm, ay, ax + gap - 3.4 * mm, ay + 1.1 * mm, ax + gap - 3.4 * mm, ay - 1.1 * mm],
                          fillColor=colors.HexColor(GRAY), strokeColor=None))
        cx, cy = x + bw / 2, top - bh - 9 * mm
        d.add(Line(cx, top - bh, cx, cy + 2.6 * mm, strokeColor=colors.HexColor(RULE), strokeWidth=0.7,
                   strokeDashArray=[1.5, 1.5]))
        r = 2.6 * mm
        d.add(Polygon([cx, cy + r, cx + r, cy, cx, cy - r, cx - r, cy], fillColor=colors.white,
                      strokeColor=colors.HexColor(BLUE), strokeWidth=0.9))
        for j, t in enumerate(gate):
            d.add(String(cx, cy - 7 * mm - j * 4 * mm, t, fontName="SC", fontSize=7.2,
                         fillColor=colors.HexColor(GRAY), textAnchor="middle"))
    d.add(String(0, 1.5 * mm, "◇ = 进入下一阶段的条件。每个阶段 Harness、现场规则、人员培训三件事同步推进，周期待了解现场后确定",
                 fontName="SC", fontSize=6.8, fillColor=colors.HexColor(GRAY)))
    return d


def write_roadmap_png(path):
    d = roadmap_drawing()
    with tempfile.TemporaryDirectory() as tmp:
        pdf = Path(tmp) / "roadmap.pdf"
        renderPDF.drawToFile(d, str(pdf))
        subprocess.run(["pdftoppm", "-png", "-r", "300", "-singlefile", str(pdf), str(path.with_suffix(""))],
                       check=True)


# ---------- PDF ----------

class HeadingRule(Flowable):
    def __init__(self, text, style):
        super().__init__()
        self.para = Paragraph(text, style)

    def wrap(self, aw, ah):
        self.aw = aw
        _, self.ph = self.para.wrap(aw, ah)
        return aw, self.ph + 3

    def draw(self):
        self.para.drawOn(self.canv, 0, 3)
        self.canv.setStrokeColor(colors.HexColor(RULE))
        self.canv.setLineWidth(0.7)
        self.canv.line(0, 0, self.aw, 0)


def build_pdf(blocks, out):
    st = {
        "title": ParagraphStyle("title", fontName="SC-Bold", fontSize=16.5, leading=23, textColor=NAVY, spaceAfter=3),
        "meta": ParagraphStyle("meta", fontName="SC", fontSize=9.5, leading=14, textColor=GRAY, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName="SC-Bold", fontSize=13.5, leading=19, textColor=NAVY),
        "h3": ParagraphStyle("h3", fontName="SC-Medium", fontSize=11, leading=16, textColor=BLUE,
                             spaceBefore=7, spaceAfter=3),
        "body": ParagraphStyle("body", fontName="SC", fontSize=9.6, leading=15.5, textColor=TEXT, spaceAfter=4),
        "cell": ParagraphStyle("cell", fontName="SC", fontSize=8.6, leading=13, textColor=TEXT),
        "cellh": ParagraphStyle("cellh", fontName="SC-Bold", fontSize=8.6, leading=13, textColor=NAVY),
        "cap": ParagraphStyle("cap", fontName="SC", fontSize=8, leading=12, textColor=GRAY, alignment=TA_CENTER),
    }
    for style in st.values():
        style.wordWrap = "CJK"
    story, first_p = [], True
    for b in blocks:
        kind = b[0]
        if kind == "h" and b[1] == 1:
            story.append(Paragraph(esc(b[2]), st["title"]))
        elif kind == "h" and b[1] == 2:
            story += [Spacer(1, 9), HeadingRule(esc(b[2]), st["h2"]), Spacer(1, 6)]
        elif kind == "h":
            story.append(Paragraph(esc(b[2]), st["h3"]))
        elif kind == "p":
            story.append(Paragraph(to_rl(b[1]), st["meta"] if first_p else st["body"]))
            first_p = False
        elif kind == "list":
            ordered, items = b[1], b[2]
            story.append(ListFlowable(
                [ListItem(Paragraph(to_rl(t), st["body"]), leftIndent=13,
                          value=(n + 1) if ordered else None) for n, t in enumerate(items)],
                bulletType="1" if ordered else "bullet", start="1" if ordered else "·",
                bulletFontName="SC-Bold" if ordered else "SC", bulletFontSize=9,
                bulletColor=colors.HexColor(BLUE), leftIndent=13, bulletDedent=11))
        elif kind == "table":
            rows = b[1]
            ncol = len(rows[0])
            widths = {2: [0.36, 0.64], 3: [0.18, 0.47, 0.35]}[ncol]
            avail = A4[0] - 40 * mm
            data = [[Paragraph(to_rl(c), st["cellh"] if r == 0 else st["cell"]) for c in row]
                    for r, row in enumerate(rows)]
            t = Table(data, colWidths=[w * avail for w in widths], repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(TABLE_HEAD)),
                ("LINEABOVE", (0, 0), (-1, 0), 0.9, colors.HexColor(NAVY)),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor(TABLE_LINE)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]))
            story += [Spacer(1, 2), t, Spacer(1, 8)]
        elif kind == "img":
            story.append(KeepTogether([Spacer(1, 4), roadmap_drawing(A4[0] - 40 * mm), Spacer(1, 5),
                                       Paragraph(esc(b[1]), st["cap"]), Spacer(1, 8)]))

    def decorate(canv, doc):
        canv.saveState()
        canv.setFont("SC-Bold", 7.5)
        canv.setFillColor(colors.HexColor(NAVY))
        canv.drawString(20 * mm, A4[1] - 12 * mm, RUNNING_HEAD)
        canv.setFont("SC", 8)
        canv.setFillColor(colors.HexColor(GRAY))
        canv.drawCentredString(A4[0] / 2, 10 * mm, f"第 {doc.page} 页")
        canv.restoreState()

    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=17 * mm,
                            title=blocks[0][2], author="鄢澳勇")
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)


# ---------- DOCX ----------

def build_docx(blocks, out, png):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    def set_font(run, size, color, bold=False):
        run.font.name = "Arial"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
        run.bold = bold

    def add_runs(par, text, size=10, color=TEXT):
        for s, kind, href in inline_runs(text):
            if kind == "link":
                rid = par.part.relate_to(href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                         is_external=True)
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), rid)
                r = par.add_run(s)
                set_font(r, size, BLUE)
                r.underline = True
                link.append(r._element)
                par._p.append(link)
            else:
                r = par.add_run(s)
                set_font(r, size, NAVY if kind == "bold" else GRAY if kind == "code" else color, kind == "bold")

    def bottom_border(par):
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in {"w:val": "single", "w:sz": "6", "w:space": "2", "w:color": RULE.lstrip("#")}.items():
            b.set(qn(k), v)
        bdr.append(b)
        pPr.append(bdr)

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color.lstrip("#"))
        tcPr.append(shd)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Mm(20))
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(17)
    hp = sec.header.paragraphs[0]
    set_font(hp.add_run(RUNNING_HEAD), 7.5, NAVY, True)

    first_p = True
    for b in blocks:
        kind = b[0]
        if kind == "h":
            p = doc.add_paragraph()
            size, color = {1: (19, NAVY), 2: (13.5, NAVY), 3: (11, BLUE)}[b[1]]
            set_font(p.add_run(b[2]), size, color, b[1] < 3)
            p.paragraph_format.space_before = Pt({1: 0, 2: 12, 3: 6}[b[1]])
            p.paragraph_format.space_after = Pt(4)
            if b[1] == 2:
                bottom_border(p)
        elif kind == "p":
            p = doc.add_paragraph()
            add_runs(p, b[1], 9.5 if first_p else 10, GRAY if first_p else TEXT)
            p.paragraph_format.line_spacing = 1.45
            first_p = False
        elif kind == "list":
            for n, t in enumerate(b[2]):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Mm(6)
                p.paragraph_format.first_line_indent = Mm(-4)
                p.paragraph_format.line_spacing = 1.45
                set_font(p.add_run(f"{n + 1}  " if b[1] else "·  "), 10, BLUE, True)
                add_runs(p, t)
        elif kind == "table":
            rows = b[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    cell = t.cell(r, c)
                    cell.paragraphs[0].text = ""
                    if r == 0:
                        shade(cell, TABLE_HEAD)
                        set_font(cell.paragraphs[0].add_run(text), 9, NAVY, True)
                    else:
                        add_runs(cell.paragraphs[0], text, 9)
            doc.add_paragraph()
        elif kind == "img":
            doc.add_picture(str(png), width=Mm(170))
            cap = doc.add_paragraph()
            set_font(cap.add_run(b[1]), 8, GRAY)
    doc.core_properties.author = "鄢澳勇"
    doc.core_properties.title = blocks[0][2]
    doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fonts", required=True, type=Path)
    args = ap.parse_args()
    for name, file in (("SC", "Regular"), ("SC-Medium", "Medium"), ("SC-Bold", "Bold")):
        pdfmetrics.registerFont(TTFont(name, str(args.fonts / f"NotoSansSC-{file}.ttf")))
    pdfmetrics.registerFontFamily("SC", normal="SC", bold="SC-Bold", italic="SC", boldItalic="SC-Bold")

    blocks = parse_blocks((HERE / "portfolio.md").read_text(encoding="utf-8"))
    png = HERE / "figures" / "roadmap.png"
    png.parent.mkdir(exist_ok=True)
    write_roadmap_png(png)
    build_pdf(blocks, HERE / "portfolio.pdf")
    build_docx(blocks, HERE / "portfolio.docx", png)


if __name__ == "__main__":
    main()
